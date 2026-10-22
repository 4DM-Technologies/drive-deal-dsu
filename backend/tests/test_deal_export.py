import asyncio
from datetime import UTC, datetime
from decimal import Decimal
from io import BytesIO

from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import update

from main import app
from src.database import SessionFactory
from src.repositories.schema import BuyerRequest
from src.services.deal_export import XLSX_MEDIA_TYPE, build_deals_workbook
from tests.demo_data import IDS

BUYER = "rahul@drivedeal.demo"
DEALER = "naveen@naveemotors.demo"


def login(client: TestClient, email: str) -> dict[str, str]:
    response = client.post("/api/v1/auth/login", json={"email": email, "password": "demo1234"})
    assert response.status_code == 200, response.text
    return {"Authorization": f"Bearer {response.json()['access_token']}"}


def set_buyer_area(request_id: str, area: str) -> None:
    async def _update() -> None:
        async with SessionFactory() as session:
            await session.execute(update(BuyerRequest).where(BuyerRequest.id == request_id).values(buyer_area=area))
            await session.commit()

    asyncio.run(_update())


def export_sheet(client: TestClient, headers: dict[str, str]):
    response = client.get("/api/v1/deals/export", headers=headers)
    assert response.status_code == 200, response.text
    assert response.headers["content-type"] == XLSX_MEDIA_TYPE
    assert response.headers["content-disposition"].startswith('attachment; filename="deals-')
    assert response.headers["cache-control"] == "no-store"
    return load_workbook(BytesIO(response.content))["Deals"]


def test_dealer_export_matches_their_deals_in_the_database() -> None:
    with TestClient(app) as client:
        dealer = login(client, DEALER)
        deals = client.get("/api/v1/deals", headers=dealer).json()
        assert deals, "the demo dealer should have accepted deals"
        sheet = export_sheet(client, dealer)

        assert sheet["A1"].value == "Deals in progress"
        assert f"{len(deals)} deal" in sheet["A2"].value
        assert [cell.value for cell in sheet[4]] == [
            "Vehicle request",
            "Year from",
            "Year to",
            "Body type",
            "Buyer area",
            "Out-the-door price",
        ]
        for row, deal in enumerate(deals, start=5):
            assert [cell.value for cell in sheet[row]] == [
                f"{deal['brand_name']} {deal['model']}",
                deal["year_min"],
                deal["year_max"],
                deal["body_type"] or "",
                deal["buyer_area"] or "",
                float(Decimal(deal["final_price"])),
            ]
            assert sheet.cell(row=row, column=6).number_format == '"$"#,##0.00'
        total_row = 5 + len(deals)
        assert sheet.cell(row=total_row, column=1).value == "Total"
        assert sheet.cell(row=total_row, column=6).value == f"=SUM(F5:F{total_row - 1})"
        assert sheet.auto_filter.ref == f"A4:F{total_row - 1}"


def test_buyer_typed_text_is_written_as_text_not_a_formula() -> None:
    formula = '=HYPERLINK("https://example.com","Click")'
    set_buyer_area(IDS["jazz"], formula)
    with TestClient(app) as client:
        sheet = export_sheet(client, login(client, DEALER))
        area = next(row[4] for row in sheet.iter_rows(min_row=5) if row[4].value == formula)
        assert area.data_type == "s"


def test_only_dealers_can_export_deals() -> None:
    with TestClient(app) as client:
        assert client.get("/api/v1/deals/export").status_code == 401
        assert client.get("/api/v1/deals/export", headers=login(client, BUYER)).status_code == 403


def test_empty_export_has_headers_but_no_total_or_filter() -> None:
    sheet = load_workbook(BytesIO(build_deals_workbook([], datetime(2026, 10, 22, tzinfo=UTC))))["Deals"]
    assert sheet["A2"].value.endswith("0 deals")
    assert sheet["A4"].value == "Vehicle request"
    assert sheet["A5"].value is None
    assert sheet.auto_filter.ref is None
