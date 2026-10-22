"""Builds the dealer's "Deals in progress" spreadsheet from the database.

The workbook is made on the server from the dealer's accepted quotes, so what the browser shows (or a user edits in
it) never decides what goes into the file.
"""

from collections.abc import Sequence
from datetime import datetime
from decimal import Decimal
from io import BytesIO

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE, Cell
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.worksheet import Worksheet

XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

# The app's brand colours.
INK = "10233F"
ACCENT = "1456B8"
ACCENT_SOFT = "E8F0FC"
BAND = "F5F8FC"
RULE = "D5DEEA"
MUTED = "6F6B64"
WHITE = "FFFFFF"

# (header, width, horizontal alignment)
COLUMNS: tuple[tuple[str, int, str], ...] = (
    ("Vehicle request", 30, "left"),
    ("Year from", 12, "center"),
    ("Year to", 12, "center"),
    ("Body type", 16, "left"),
    ("Buyer area", 28, "left"),
    ("Out-the-door price", 22, "right"),
)
HEADER_ROW = 4
PRICE_COLUMN = 6
PRICE_FORMAT = '"$"#,##0.00'

THIN = Side(style="thin", color=RULE)
GRID = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def _font(**overrides: object) -> Font:
    return Font(**{"name": "Calibri", "size": 11, "color": INK, **overrides})


def _fill(color: str) -> PatternFill:
    return PatternFill(fill_type="solid", start_color=color, end_color=color)


def _alignment(horizontal: str) -> Alignment:
    return Alignment(vertical="center", horizontal=horizontal, indent=0 if horizontal == "center" else 1)


def _text(cell: Cell, value: str | None) -> None:
    """Writes buyer- or dealer-typed text as plain text.

    openpyxl stores any string that starts with "=" as a formula, so a buyer area like ``=HYPERLINK(...)`` would run
    when the dealer opens the file. Forcing the string type keeps it as the literal text. Characters XML cannot hold
    are dropped, since they would make the whole file unreadable.
    """
    cell.value = ILLEGAL_CHARACTERS_RE.sub("", value or "")
    cell.data_type = "s"


def _title_block(sheet: Worksheet, deal_count: int, exported_at: datetime) -> None:
    last = sheet.cell(row=1, column=len(COLUMNS)).column_letter
    sheet.merge_cells(f"A1:{last}1")
    sheet["A1"] = "Deals in progress"
    sheet["A1"].font = _font(size=16, bold=True)
    sheet.row_dimensions[1].height = 26
    sheet.merge_cells(f"A2:{last}2")
    plural = "deal" if deal_count == 1 else "deals"
    sheet["A2"] = (
        f"Deal&Drive  ·  Exported {exported_at:%B} {exported_at.day}, {exported_at.year}  ·  {deal_count} {plural}"
    )
    sheet["A2"].font = _font(size=10, color=MUTED)
    sheet.row_dimensions[3].height = 8


def build_deals_workbook(deals: Sequence[dict], exported_at: datetime) -> bytes:
    """``deals`` are the dealer's accepted quotes as the API serializes them (``MarketplaceService.list_deals``)."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Deals"
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = sheet.cell(row=HEADER_ROW + 1, column=1)
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.print_title_rows = f"{HEADER_ROW}:{HEADER_ROW}"
    workbook.properties.creator = "Deal&Drive"
    workbook.properties.created = exported_at

    _title_block(sheet, len(deals), exported_at)

    accent_edge = Side(style="thin", color=ACCENT)
    for column, (header, width, align) in enumerate(COLUMNS, start=1):
        cell = sheet.cell(row=HEADER_ROW, column=column, value=header)
        cell.font = _font(bold=True, color=WHITE)
        cell.fill = _fill(ACCENT)
        cell.border = Border(left=THIN, right=THIN, top=accent_edge, bottom=accent_edge)
        cell.alignment = _alignment(align)
        sheet.column_dimensions[cell.column_letter].width = width
    sheet.row_dimensions[HEADER_ROW].height = 24

    for index, deal in enumerate(deals):
        row = HEADER_ROW + 1 + index
        brand = deal.get("brand_name")
        _text(
            sheet.cell(row=row, column=1), f"{brand} {deal.get('model') or ''}".strip() if brand else "Vehicle request"
        )
        sheet.cell(row=row, column=2, value=deal.get("year_min"))
        sheet.cell(row=row, column=3, value=deal.get("year_max"))
        _text(sheet.cell(row=row, column=4), deal.get("body_type"))
        _text(sheet.cell(row=row, column=5), deal.get("buyer_area"))
        price = sheet.cell(row=row, column=PRICE_COLUMN, value=Decimal(str(deal["final_price"])))
        price.number_format = PRICE_FORMAT
        for column, (_, _, align) in enumerate(COLUMNS, start=1):
            cell = sheet.cell(row=row, column=column)
            cell.font = _font(bold=column == 1)
            cell.border = GRID
            cell.alignment = _alignment(align)
            if index % 2 == 1:
                cell.fill = _fill(BAND)
        sheet.row_dimensions[row].height = 21

    if deals:
        last_row = HEADER_ROW + len(deals)
        price_letter = sheet.cell(row=HEADER_ROW, column=PRICE_COLUMN).column_letter
        sheet.auto_filter.ref = f"A{HEADER_ROW}:{price_letter}{last_row}"
        total_row = last_row + 1
        sheet.cell(row=total_row, column=1, value="Total")
        total = sheet.cell(
            row=total_row, column=PRICE_COLUMN, value=f"=SUM({price_letter}{HEADER_ROW + 1}:{price_letter}{last_row})"
        )
        total.number_format = PRICE_FORMAT
        for column in range(1, len(COLUMNS) + 1):
            cell = sheet.cell(row=total_row, column=column)
            cell.font = _font(bold=True)
            cell.fill = _fill(ACCENT_SOFT)
            cell.border = Border(left=THIN, right=THIN, bottom=THIN, top=Side(style="medium", color=ACCENT))
            cell.alignment = _alignment("right" if column == PRICE_COLUMN else "left")
        sheet.row_dimensions[total_row].height = 23

    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
